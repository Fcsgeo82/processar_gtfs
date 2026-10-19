# -*- coding: utf-8 -*-
"""
1.3.1 — Conferir partidas geradas pelo script 1.3 contra a tabela de entrada.

Conferência INDEPENDENTE (não confia no _resumo.txt nem no relatorio_bandas.csv):
rele a tabela de contagens de entrada e os arquivos horarios_*.csv gravados no
disco, reexpande as bandas e compara as contagens de partidas por
(servico, sentido, faixa). Opcionalmente confere o Excel gerado (uma aba por
linha, blocos por quadro).

Anomalias verificadas:
  - FALTANDO : grupo com partidas>0 na entrada sem arquivo de quadro correspondente
  - SOBRANDO : grupo com partidas=0 na entrada que possui arquivo de quadro
  - ORFAO    : arquivo de quadro sem grupo correspondente na entrada
  - BANDA    : violação em banda do disco (headway < 1, não divisível, fora da faixa)
  - CELULA   : divergência de contagem por (servico, sentido, faixa)
  - EXCEL    : divergência entre o Excel e a expansão dos CSVs

Saídas na pasta conferida: relatorio_conferencia_partidas.csv + _conferencia_resumo.txt.
Exit code 0 = tudo OK; 1 = pelo menos uma anomalia.

Uso:
  python codigos_py/1.3.1_conferir_partidas_geradas.py --config config.json
  python codigos_py/1.3.1_conferir_partidas_geradas.py   # usa defaults/hardcoded abaixo
"""
import os
import re
import json
import unicodedata
import argparse
from pathlib import Path
import pandas as pd

# ==============================================================================
# LEITURA DE CONFIGURAÇÃO (Linha de Comando / Streamlit / JSON)
# ==============================================================================
_parser = argparse.ArgumentParser(
    description="Confere partidas dos quadros gerados pelo 1.3 contra a tabela de contagens de entrada."
)
_parser.add_argument('--config', type=str, default='', help='Caminho para arquivo JSON de configuração')
_args, _ = _parser.parse_known_args()
_config = {}
if _args.config:
    if os.path.isfile(_args.config):
        with open(_args.config, 'r', encoding='utf-8-sig') as _f:
            _config = json.load(_f)
    else:
        try:
            _config = json.loads(_args.config)
        except Exception:
            with open(_args.config, 'r', encoding='utf-8-sig') as _f:
                _config = json.load(_f)

# Parâmetros principais
BASE_RESULTADOS = Path(_config.get("BASE_RESULTADOS", "C:/R_SMTR/resultados"))

ano_gtfs = str(_config.get("ano_gtfs", "2030"))
mes_gtfs = str(_config.get("mes_gtfs", "10")).zfill(2)
estudo_gtfs = str(_config.get("estudo_gtfs", "01"))

# Caminho da tabela de contagens usada na geração (formato wide ou longo)
caminho_partidas = Path(_config.get("caminho_partidas", "partidas_pf_sppo_como_du.csv"))

# service_id dos arquivos gerados (usado no nome dos CSVs e do Excel)
service_id = str(_config.get("service_id", "F_REG"))

# Tratamento de contagem '-' na coluna de partidas (igual ao 1.3)
traco_ausente = str(_config.get("traco_ausente", "zerar"))
if traco_ausente not in ("zerar", "erro"):
    raise ValueError(f"traco_ausente inválido: {traco_ausente!r} (use 'zerar' ou 'erro')")

# Linhas a rodar (vazio ou ['todas'] = todas) — deve refletir o filtro usado na geração
linhas_rodar = _config.get("linhas_rodar", [])
if isinstance(linhas_rodar, str):
    linhas_rodar = [x.strip() for x in linhas_rodar.split(',') if x.strip()]
linhas_rodar_set = {str(x).upper() for x in linhas_rodar} if linhas_rodar and ('todas' not in [str(x).lower() for x in linhas_rodar]) else set()

# Pasta de saída = pasta conferida (onde estão os horarios_*.csv)
pasta_saida_config = _config.get("pasta_saida", "")
if pasta_saida_config:
    pasta_saida = Path(pasta_saida_config)
else:
    pasta_saida = BASE_RESULTADOS / f"quadro_horario_extraido/{ano_gtfs}/{mes_gtfs}/qh_por_linha_partidas/{estudo_gtfs}Q"

# Confere também o Excel gerado?
verificar_excel = bool(_config.get("verificar_excel", True))
nome_excel = str(_config.get("nome_excel", f"quadros_partidas_{service_id}.xlsx"))

# ==============================================================================
# CONSTANTES E FUNÇÕES UTILITÁRIAS (mesmos critérios do script 1.3)
# ==============================================================================
FAIXAS_PADRAO = {
    '00:00-01:00': (0, 3600), '01:00-02:00': (3600, 7200), '02:00-03:00': (7200, 10800),
    '03:00-04:00': (10800, 14400), '04:00-05:00': (14400, 18000), '05:00-06:00': (18000, 21600),
    '06:00-09:00': (21600, 32400), '09:00-12:00': (32400, 43200), '12:00-15:00': (43200, 54000),
    '15:00-18:00': (54000, 64800), '18:00-21:00': (64800, 75600), '21:00-22:00': (75600, 79200),
    '22:00-23:00': (79200, 82800), '23:00-24:00': (82800, 86400),
}

SENTIDOS_VALIDOS = {'Ida', 'Volta', 'Circular'}


def sanitize_filename(name: str) -> str:
    """Substitui caracteres inválidos para nomes de arquivos no sistema operacional."""
    if not name:
        return "SEM_NOME"
    sanitized = re.sub(r'[\\/*?:"<>|]', '_', str(name))
    sanitized = re.sub(r'\s+', ' ', sanitized).strip('. ')
    return sanitized if sanitized else "SEM_NOME"


def horario_to_seconds(hk_str) -> int:
    """Converte horário HH:MM:SS ou HH:MM para total de segundos (suporta H >= 24)."""
    if pd.isna(hk_str) or not str(hk_str).strip():
        return -1
    parts = str(hk_str).strip().split(':')
    if len(parts) < 2:
        return -1
    try:
        h = int(parts[0])
        m = int(parts[1])
        s = int(parts[2]) if len(parts) >= 3 else 0
        return h * 3600 + m * 60 + s
    except (ValueError, TypeError):
        return -1


def normalizar_coluna(nome: str) -> str:
    """Normaliza nomes de coluna: minúsculas, sem acentos e sem espaços nas bordas."""
    nome = unicodedata.normalize('NFKD', str(nome)).encode('ascii', 'ignore').decode('ascii')
    return re.sub(r'\s+', '_', nome.strip().lower())


def derivar_headsign(vista: str, sentido: str) -> str:
    """Deriva o trip_headsign a partir da vista ('Origem - Destino') e do sentido."""
    if sentido == 'Circular':
        return 'Circular'
    partes = str(vista).split(' - ')
    if len(partes) != 2 or not partes[0].strip() or not partes[1].strip():
        raise ValueError(f"vista {vista!r} não tem exatamente 2 partes separadas por ' - ' (sentido {sentido})")
    origem, destino = partes[0].strip(), partes[1].strip()
    return destino if sentido == 'Ida' else origem


def faixa_de_inicio(segundos: int):
    """Retorna o rótulo da faixa padrão que contém o instante (segundos) ou None."""
    for rotulo, (ini, fim) in FAIXAS_PADRAO.items():
        if ini <= segundos < fim:
            return rotulo
    return None


# ==============================================================================
# EXECUÇÃO PRINCIPAL
# ==============================================================================
print("==========================================================")
print("CONFERÊNCIA DE PARTIDAS — QUADROS GERADOS PELO SCRIPT 1.3")
print(f"Tabela de entrada : {caminho_partidas}")
print(f"Pasta conferida   : {pasta_saida}")
print(f"service_id        : {service_id}")
print(f"Conferir Excel    : {'sim' if verificar_excel else 'não'}")
print("==========================================================\n")

if not caminho_partidas.exists():
    raise FileNotFoundError(f"Tabela de partidas não encontrada: {caminho_partidas}")
if not pasta_saida.exists():
    raise FileNotFoundError(f"Pasta de quadros não encontrada: {pasta_saida}")

# 1. Leitura da tabela de entrada (wide com melt / longo — igual ao 1.3)
try:
    df = pd.read_csv(caminho_partidas, dtype=str, encoding='utf-8-sig')
except UnicodeDecodeError:
    df = pd.read_csv(caminho_partidas, dtype=str, encoding='latin-1')

df.columns = [normalizar_coluna(c) for c in df.columns]

if not ('faixa' in df.columns and 'partidas_faixa' in df.columns):
    mapa_faixa = {}
    for f in FAIXAS_PADRAO:
        for c in df.columns:
            if re.sub(r'[\s_]+', '', c) == f:
                mapa_faixa[c] = f
                break
    if not mapa_faixa:
        raise ValueError(
            "Tabela de entrada não está nem no formato longo "
            "(colunas 'faixa' e 'partidas_faixa') nem no formato wide "
            f"(colunas uma por faixa, ex.: '{next(iter(FAIXAS_PADRAO))}'). "
            f"Colunas lidas: {list(df.columns)}"
        )
    df = df.rename(columns=mapa_faixa)
    id_vars = [c for c in df.columns if c not in set(mapa_faixa.values())]
    df = df.melt(
        id_vars=id_vars,
        value_vars=list(mapa_faixa.values()),
        var_name='faixa',
        value_name='partidas_faixa',
    )
    print(f"1. Formato detectado: wide ({len(mapa_faixa)} faixas derretidas)")

obrigatorias = ['servico', 'vista', 'sentido', 'faixa', 'partidas_faixa']
faltantes = [c for c in obrigatorias if c not in df.columns]
if faltantes:
    raise ValueError(f"Colunas obrigatórias ausentes na tabela de entrada: {faltantes}. Colunas lidas: {list(df.columns)}")

# Contagens (mesmo critério do 1.3: '-'/'vazio' -> 0 com aviso ou erro)
valores = df['partidas_faixa'].fillna('').astype(str).str.strip()
mascara_traco = ~valores.str.fullmatch(r'\d+')
n_tracos = int(mascara_traco.sum())
if n_tracos:
    if traco_ausente == 'erro':
        raise ValueError(f"{n_tracos} valores não numéricos na coluna partidas_faixa (traco_ausente='erro').")
    print(f"   Aviso: {n_tracos} valores '-'/'vazios' em partidas_faixa tratados como 0.")
    valores = valores.where(~mascara_traco, '0')
df['n_entrada'] = pd.to_numeric(valores, errors='coerce')
if df['n_entrada'].isna().any():
    raise ValueError("Contagens não conversíveis em partidas_faixa.")
df['n_entrada'] = df['n_entrada'].astype(int)

# Filtro de linhas (deve ser idêntico ao usado na geração)
if linhas_rodar_set:
    df = df[df['servico'].str.upper().isin(linhas_rodar_set)].copy()
    print(f"   Filtro de linhas aplicado: {df['servico'].nunique()} serviços restantes")

# Esperados: por (servico, sentido) -> headsign, vista e total
esperados = {}
for (servico, sentido), grupo in df.groupby(['servico', 'sentido'], sort=True):
    vista = grupo['vista'].iloc[0]
    try:
        headsign = derivar_headsign(vista, sentido)
    except ValueError as e:
        raise ValueError(f"Grupo ({servico}, {sentido}): {e}")
    esperados[(servico, sentido)] = {
        'headsign': headsign,
        'vista': vista,
        'faixas': dict(zip(grupo['faixa'], grupo['n_entrada'])),
        'total': int(grupo['n_entrada'].sum()),
    }

total_entrada = sum(e['total'] for e in esperados.values())
grupos_com_partidas = sum(1 for e in esperados.values() if e['total'] > 0)
print(f"2. Entrada: {len(df):,} células, {len(esperados)} grupos, {grupos_com_partidas} com partidas, total {total_entrada:,}")

# Mapeia headsign esperado -> sentido por serviço (para identificar grupos dos CSVs)
headsign_para_sentido = {}
for (servico, sentido), info in esperados.items():
    headsign_para_sentido[(servico, info['headsign'])] = sentido

anomalias = []   # linhas de relatório não-OK
celulas = []     # linhas de relatório por célula (servico, sentido, faixa)
expansao_por_arquivo = {}  # nome_arquivo -> total expandido
expansao_por_grupo = {}    # (servico, sentido) -> {faixa: n}

# 3. Varre os arquivos horarios_*.csv no disco
arquivos_quadro = sorted(pasta_saida.glob(f"horarios_*_{service_id}.csv"))
print(f"3. Arquivos de quadro no disco: {len(arquivos_quadro):,}")

arquivos_lidos = {}  # (servico, sentido) -> nome do arquivo

for caminho_arq in arquivos_quadro:
    df_arq = pd.read_csv(caminho_arq, dtype=str, encoding='utf-8')
    if df_arq.empty:
        anomalias.append({
            'tipo': 'ORFAO', 'servico': '', 'sentido': '', 'headsign': '', 'faixa': '',
            'n_entrada': '', 'n_quadro': '', 'detalhe': f'arquivo vazio: {caminho_arq.name}', 'status': 'ANOMALIA',
        })
        continue

    servicos_no_arq = set(df_arq['trip_short_name'].dropna().unique())
    headsigns_no_arq = set(df_arq['trip_headsign'].dropna().unique())
    if len(servicos_no_arq) != 1 or len(headsigns_no_arq) != 1:
        anomalias.append({
            'tipo': 'ORFAO', 'servico': '|'.join(sorted(servicos_no_arq)), 'sentido': '',
            'headsign': '|'.join(sorted(headsigns_no_arq)), 'faixa': '',
            'n_entrada': '', 'n_quadro': '', 'detalhe': f'arquivo com múltiplos serviços/headsigns: {caminho_arq.name}',
            'status': 'ANOMALIA',
        })
        continue

    servico = servicos_no_arq.pop()
    headsign = headsigns_no_arq.pop()
    sentido = headsign_para_sentido.get((servico, headsign))

    if sentido is None:
        # Não corresponde a nenhum grupo da entrada (nem por headsign derivado)
        sens = [s for (sv, s) in esperados if sv == servico]
        anomalias.append({
            'tipo': 'ORFAO', 'servico': servico, 'sentido': '/'.join(sens) or '?',
            'headsign': headsign, 'faixa': '',
            'n_entrada': '', 'n_quadro': int(len(df_arq)), 'detalhe': f'arquivo sem grupo correspondente na entrada: {caminho_arq.name}',
            'status': 'ANOMALIA',
        })
        continue

    info = esperados[(servico, sentido)]
    if info['total'] == 0:
        anomalias.append({
            'tipo': 'SOBRANDO', 'servico': servico, 'sentido': sentido, 'headsign': headsign, 'faixa': '',
            'n_entrada': 0, 'n_quadro': int(len(df_arq)), 'detalhe': f'grupo com 0 partidas na entrada possui arquivo: {caminho_arq.name}',
            'status': 'ANOMALIA',
        })

    if (servico, sentido) in arquivos_lidos:
        anomalias.append({
            'tipo': 'ORFAO', 'servico': servico, 'sentido': sentido, 'headsign': headsign, 'faixa': '',
            'n_entrada': '', 'n_quadro': '', 'detalhe': f'grupo com mais de um arquivo (duplicado de {arquivos_lidos[(servico, sentido)]}): {caminho_arq.name}',
            'status': 'ANOMALIA',
        })
        continue
    arquivos_lidos[(servico, sentido)] = caminho_arq.name

    # Reexpande as bandas e valida limites
    por_faixa = {}
    for _, banda in df_arq.iterrows():
        ini_b = horario_to_seconds(banda['start_time'])
        fim_b = horario_to_seconds(banda['end_time'])
        # headway_secs é gravado como total de segundos (ex.: 3600); aceita também HH:MM:SS
        hb_str = str(banda['headway_secs']).strip()
        if re.fullmatch(r'\d+', hb_str):
            h_b = int(hb_str)
        else:
            h_b = horario_to_seconds(hb_str)
        detalhe_problema = None
        if h_b < 1:
            detalhe_problema = f'headway < 1: {banda["headway_secs"]}'
        elif fim_b <= ini_b:
            detalhe_problema = f'end_time <= start_time: {banda["start_time"]} -> {banda["end_time"]}'
        elif (fim_b - ini_b) % h_b != 0:
            detalhe_problema = f'duração não divisível por headway: {fim_b - ini_b}s / {h_b}s'
        faixa = faixa_de_inicio(ini_b)
        if detalhe_problema is None and faixa is None:
            detalhe_problema = f'start_time fora das 14 faixas padrão: {banda["start_time"]}'
        if detalhe_problema is None and fim_b > FAIXAS_PADRAO[faixa][1]:
            fim_faixa = FAIXAS_PADRAO[faixa][1]
            detalhe_problema = f'end_time vaza da faixa {faixa}: {banda["end_time"]} > {fim_faixa // 3600:02d}:{(fim_faixa % 3600) // 60:02d}:00'

        if detalhe_problema is not None:
            anomalias.append({
                'tipo': 'BANDA', 'servico': servico, 'sentido': sentido, 'headsign': headsign,
                'faixa': faixa or '', 'n_entrada': '', 'n_quadro': '',
                'detalhe': f'{caminho_arq.name}: {detalhe_problema}', 'status': 'ANOMALIA',
            })
            continue

        n_banda = (fim_b - ini_b) // h_b
        por_faixa[faixa] = por_faixa.get(faixa, 0) + n_banda

    expansao_por_grupo[(servico, sentido)] = por_faixa
    expansao_por_arquivo[caminho_arq.name] = sum(por_faixa.values())

# 4. Grupos esperados com partidas>0 sem arquivo -> FALTANDO
for (servico, sentido), info in esperados.items():
    if info['total'] > 0 and (servico, sentido) not in arquivos_lidos:
        anomalias.append({
            'tipo': 'FALTANDO', 'servico': servico, 'sentido': sentido, 'headsign': info['headsign'],
            'faixa': '', 'n_entrada': info['total'], 'n_quadro': 0,
            'detalhe': 'grupo com partidas>0 sem arquivo de quadro', 'status': 'ANOMALIA',
        })

# 5. Comparação célula a célula (servico, sentido, faixa)
n_ok = 0
n_divergente = 0
todos_keys = set()
for (servico, sentido), info in esperados.items():
    for faixa in info['faixas']:
        todos_keys.add((servico, sentido, faixa))
for (servico, sentido), por_faixa in expansao_por_grupo.items():
    for faixa in por_faixa:
        todos_keys.add((servico, sentido, faixa))

for (servico, sentido, faixa) in sorted(todos_keys):
    n_ent = esperados.get((servico, sentido), {}).get('faixas', {}).get(faixa, 0)
    n_quad = expansao_por_grupo.get((servico, sentido), {}).get(faixa, 0)
    headsign = esperados.get((servico, sentido), {}).get('headsign', '')
    status = 'OK' if n_ent == n_quad else 'DIVERGENTE'
    if status == 'OK':
        n_ok += 1
    else:
        n_divergente += 1
        anomalias.append({
            'tipo': 'CELULA', 'servico': servico, 'sentido': sentido, 'headsign': headsign,
            'faixa': faixa, 'n_entrada': n_ent, 'n_quadro': n_quad,
            'detalhe': f'delta {n_quad - n_ent:+d}', 'status': 'ANOMALIA',
        })
    celulas.append({
        'tipo': 'CELULA', 'servico': servico, 'sentido': sentido, 'headsign': headsign,
        'faixa': faixa, 'n_entrada': n_ent, 'n_quadro': n_quad,
        'detalhe': '' if status == 'OK' else f'delta {n_quad - n_ent:+d}', 'status': status,
    })

# 6. Conferência do Excel (blocos por quadro vs expansão dos CSVs)
n_excel_ok = 0
n_excel_div = 0
if verificar_excel:
    caminho_excel = pasta_saida / nome_excel
    if not caminho_excel.exists():
        anomalias.append({
            'tipo': 'EXCEL', 'servico': '', 'sentido': '', 'headsign': '', 'faixa': '',
            'n_entrada': '', 'n_quadro': '', 'detalhe': f'arquivo não encontrado: {nome_excel}',
            'status': 'ANOMALIA',
        })
    else:
        from openpyxl import load_workbook
        wb = load_workbook(caminho_excel, read_only=True, data_only=True)
        contagem_excel = {}  # (servico, headsign) -> n partidas no bloco
        for ws in wb.worksheets:
            servico_bloco = None
            headsign_bloco = None
            em_tabela = False
            for row in ws.iter_rows(values_only=True):
                c0 = row[0] if len(row) > 0 else None
                if c0 == 'Linha':
                    servico_bloco = str(row[1]) if len(row) > 1 and row[1] is not None else None
                    headsign_bloco = str(row[3]) if len(row) > 3 and row[3] is not None else None
                    em_tabela = False
                    continue
                if c0 == 'Partida':
                    em_tabela = True
                    continue
                if em_tabela and isinstance(c0, (int, float)) and servico_bloco is not None:
                    chave = (servico_bloco, headsign_bloco)
                    contagem_excel[chave] = contagem_excel.get(chave, 0) + 1
        wb.close()

        # Compara cada bloco do Excel com a expansão do CSV correspondente
        for (xl_servico, xl_headsign), n_xl in sorted(contagem_excel.items()):
            sentido = headsign_para_sentido.get((xl_servico, xl_headsign))
            if sentido is None:
                anomalias.append({
                    'tipo': 'EXCEL', 'servico': xl_servico, 'sentido': '', 'headsign': xl_headsign,
                    'faixa': '', 'n_entrada': '', 'n_quadro': n_xl,
                    'detalhe': 'bloco do Excel sem grupo correspondente na entrada', 'status': 'ANOMALIA',
                })
                n_excel_div += 1
                continue
            n_csv = sum(expansao_por_grupo.get((xl_servico, sentido), {}).values())
            if n_xl == n_csv:
                n_excel_ok += 1
            else:
                n_excel_div += 1
                anomalias.append({
                    'tipo': 'EXCEL', 'servico': xl_servico, 'sentido': sentido, 'headsign': xl_headsign,
                    'faixa': '', 'n_entrada': n_csv, 'n_quadro': n_xl,
                    'detalhe': f'Excel != expansão dos CSVs (delta {n_xl - n_csv:+d})', 'status': 'ANOMALIA',
                })

        total_excel = sum(contagem_excel.values())
        total_csv = sum(expansao_por_arquivo.values())
        if total_excel != total_csv:
            anomalias.append({
                'tipo': 'EXCEL', 'servico': '', 'sentido': '', 'headsign': '', 'faixa': '',
                'n_entrada': total_csv, 'n_quadro': total_excel,
                'detalhe': f'total do Excel != total da expansão dos CSVs (delta {total_excel - total_csv:+d})',
                'status': 'ANOMALIA',
            })

# 7. Relatórios
total_quadro = sum(expansao_por_arquivo.values())
n_anomalias = len(anomalias)

relatorio = pd.DataFrame(celulas + anomalias)
colunas_rel = ['tipo', 'servico', 'sentido', 'headsign', 'faixa', 'n_entrada', 'n_quadro', 'detalhe', 'status']
relatorio = relatorio[colunas_rel] if not relatorio.empty else pd.DataFrame(columns=colunas_rel)
caminho_relatorio = pasta_saida / 'relatorio_conferencia_partidas.csv'
relatorio.to_csv(caminho_relatorio, index=False, encoding='utf-8')

resumo = [
    "RESUMO — 1.3.1 Conferência de partidas geradas",
    f"Tabela de entrada       : {caminho_partidas}",
    f"Pasta conferida         : {pasta_saida}",
    f"Grupos na entrada       : {len(esperados)} ({grupos_com_partidas} com partidas)",
    f"Partidas na entrada     : {total_entrada:,}",
    f"Arquivos no disco       : {len(arquivos_lidos)}",
    f"Partidas expandidas     : {total_quadro:,}",
    f"Células conferidas      : {len(celulas):,} (OK {n_ok:,} | divergentes {n_divergente:,})",
]
if verificar_excel:
    resumo.append(f"Excel blocos            : OK {n_excel_ok} | divergentes {n_excel_div}")
resumo.append(f"Anomalias               : {n_anomalias}")
resumo.append(f"Resultado               : {'OK — contagens conferidas' if n_anomalias == 0 else 'FALHOU — ver relatório'}")

caminho_resumo = pasta_saida / '_conferencia_resumo.txt'
caminho_resumo.write_text('\n'.join(resumo) + '\n', encoding='utf-8')

print("\n----------------------------------------------------------")
for linha in resumo:
    print(linha)
print("----------------------------------------------------------")
print(f"Relatório: {caminho_relatorio}")
print(f"Resumo   : {caminho_resumo}")

if n_anomalias:
    print("\n!!! CONFERÊNCIA FALHOU — primeiras anomalias:")
    for a in anomalias[:15]:
        print(f"    - [{a['tipo']}] {a['servico']}/{a['sentido']} {a['faixa']}: {a['detalhe']}")
    raise SystemExit(1)

print("\n==========================================================")
print("CONFERÊNCIA CONCLUÍDA — CONTAGENS CONFERIDAS!")
print("==========================================================")
