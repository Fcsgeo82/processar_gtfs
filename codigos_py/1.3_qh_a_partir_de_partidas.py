import os
import re
import json
import unicodedata
import argparse
from pathlib import Path
import pandas as pd

# ==============================================================================
# LEITURA DE CONFIGURAÇÃO (Linha de Comando / Streamlit / JSON)
# ==============================================================================
_parser = argparse.ArgumentParser(
    description="Gera quadros de horário (bandas de frequência, formato do script 1) a partir de tabela de contagens de partidas por linha/sentido/faixa."
)
_parser.add_argument('--config', type=str, default='', help='Caminho para arquivo JSON de configuração')
_args, _ = _parser.parse_known_args()
_config = {}
if _args.config:
    if os.path.isfile(_args.config):
        with open(_args.config, 'r', encoding='utf-8-sig') as _f:
            _config = json.load(_f)
    else:
        try:
            _config = json.loads(_args.config)
        except Exception:
            with open(_args.config, 'r', encoding='utf-8-sig') as _f:
                _config = json.load(_f)

# Parâmetros principais
BASE_RESULTADOS = Path(_config.get("BASE_RESULTADOS", "C:/R_SMTR/resultados"))

ano_gtfs = str(_config.get("ano_gtfs", "2030"))
mes_gtfs = str(_config.get("mes_gtfs", "10")).zfill(2)
estudo_gtfs = str(_config.get("estudo_gtfs", "01"))

# Caminho da tabela de contagens (formato wide: uma coluna por faixa;
# ou formato longo: colunas Faixa, Partidas_Faixa)
caminho_partidas = Path(_config.get("caminho_partidas", "partidas_pf_sppo_como_du.csv"))

# service_id dos arquivos gerados
service_id = str(_config.get("service_id", "F_REG"))

# Granularidade do headway (segundos): h é arredondado para baixo nesse múltiplo
granularidade_headway = int(_config.get("granularidade_headway", 300))

# Tratamento de contagem '-' na coluna de partidas: 'zerar' (0 com aviso) ou 'erro'
traco_ausente = str(_config.get("traco_ausente", "zerar"))
if traco_ausente not in ("zerar", "erro"):
    raise ValueError(f"traco_ausente inválido: {traco_ausente!r} (use 'zerar' ou 'erro')")

# Linhas a rodar (vazio ou ['todas'] = todas)
linhas_rodar = _config.get("linhas_rodar", [])
if isinstance(linhas_rodar, str):
    linhas_rodar = [x.strip() for x in linhas_rodar.split(',') if x.strip()]
linhas_rodar_set = {str(x).upper() for x in linhas_rodar} if linhas_rodar and ('todas' not in [str(x).lower() for x in linhas_rodar]) else set()

# Pasta de saída
pasta_saida_config = _config.get("pasta_saida", "")
if pasta_saida_config:
    pasta_saida = Path(pasta_saida_config)
else:
    pasta_saida = BASE_RESULTADOS / f"quadro_horario_extraido/{ano_gtfs}/{mes_gtfs}/qh_por_linha_partidas/{estudo_gtfs}Q"

gerar_relatorio = bool(_config.get("gerar_relatorio", True))

# Exportação Excel: um workbook com uma aba por linha (servico);
# em cada aba, um bloco por quadro (sentido/headsign) com as colunas
# Partida | Horário | Intervalo. O rótulo de tipo dia é o service_id.
gerar_excel = bool(_config.get("gerar_excel", True))
nome_excel = str(_config.get("nome_excel", f"quadros_partidas_{service_id}.xlsx"))
tipo_dia_rotulo = str(_config.get("tipo_dia", service_id))

# ==============================================================================
# CONSTANTES E FUNÇÕES UTILITÁRIAS
# ==============================================================================
# As 14 faixas horárias padrão do projeto (rótulo -> início/fim em segundos)
FAIXAS_PADRAO = {
    '00:00-01:00': (0, 3600), '01:00-02:00': (3600, 7200), '02:00-03:00': (7200, 10800),
    '03:00-04:00': (10800, 14400), '04:00-05:00': (14400, 18000), '05:00-06:00': (18000, 21600),
    '06:00-09:00': (21600, 32400), '09:00-12:00': (32400, 43200), '12:00-15:00': (43200, 54000),
    '15:00-18:00': (54000, 64800), '18:00-21:00': (64800, 75600), '21:00-22:00': (75600, 79200),
    '22:00-23:00': (79200, 82800), '23:00-24:00': (82800, 86400),
}

SENTIDOS_VALIDOS = {'Ida', 'Volta', 'Circular'}


def sanitize_filename(name: str) -> str:
    """Substitui caracteres inválidos para nomes de arquivos no sistema operacional."""
    if not name:
        return "SEM_NOME"
    sanitized = re.sub(r'[\\/*?:"<>|]', '_', str(name))
    sanitized = re.sub(r'\s+', ' ', sanitized).strip('. ')
    return sanitized if sanitized else "SEM_NOME"


def horario_to_seconds(hk_str) -> int:
    """Converte horário HH:MM:SS ou HH:MM para total de segundos (suporta H >= 24)."""
    if pd.isna(hk_str) or not str(hk_str).strip():
        return -1
    parts = str(hk_str).strip().split(':')
    if len(parts) < 2:
        return -1
    try:
        h = int(parts[0])
        m = int(parts[1])
        s = int(parts[2]) if len(parts) >= 3 else 0
        return h * 3600 + m * 60 + s
    except (ValueError, TypeError):
        return -1


def seconds_to_horario(total_seconds: int) -> str:
    """Converte total de segundos de volta para string HH:MM:SS mantendo H >= 24."""
    if total_seconds < 0:
        return ""
    h = total_seconds // 3600
    m = (total_seconds % 3600) // 60
    s = total_seconds % 60
    return f"{h:02d}:{m:02d}:{s:02d}"


def normalizar_coluna(nome: str) -> str:
    """Normaliza nomes de coluna: minúsculas, sem acentos e sem espaços nas bordas."""
    nome = unicodedata.normalize('NFKD', str(nome)).encode('ascii', 'ignore').decode('ascii')
    return re.sub(r'\s+', '_', nome.strip().lower())


def derivar_headsign(vista: str, sentido: str) -> str:
    """Deriva o trip_headsign a partir da vista ('Origem - Destino') e do sentido."""
    if sentido == 'Circular':
        return 'Circular'
    partes = str(vista).split(' - ')
    if len(partes) != 2 or not partes[0].strip() or not partes[1].strip():
        raise ValueError(f"vista {vista!r} não tem exatamente 2 partes separadas por ' - ' (sentido {sentido})")
    origem, destino = partes[0].strip(), partes[1].strip()
    return destino if sentido == 'Ida' else origem


# ==============================================================================
# EXECUÇÃO PRINCIPAL
# ==============================================================================
print("==========================================================")
print("GERADOR DE QUADROS DE HORÁRIO A PARTIR DE CONTAGENS")
print(f"Tabela de Entrada : {caminho_partidas}")
print(f"Pasta de Destino  : {pasta_saida}")
print(f"service_id        : {service_id}")
print(f"Granularidade     : {granularidade_headway}s")
print("==========================================================\n")

if not caminho_partidas.exists():
    raise FileNotFoundError(f"Tabela de partidas não encontrada: {caminho_partidas}")

# 1. Leitura e normalização das colunas
try:
    df = pd.read_csv(caminho_partidas, dtype=str, encoding='utf-8-sig')
except UnicodeDecodeError:
    df = pd.read_csv(caminho_partidas, dtype=str, encoding='latin-1')

df.columns = [normalizar_coluna(c) for c in df.columns]

# Detecta o formato da tabela de entrada:
#   - longo : uma linha por (servico, sentido, faixa) com as colunas
#             faixa + partidas_faixa
#   - wide  : uma linha por (servico, sentido) e uma coluna por faixa
#             (ex.: 00:00-01:00, 01:00-02:00, ...)
# No formato wide as colunas de faixa são derretidas (melt) para o formato
# longo, tornando o restante do processamento idêntico nos dois casos.
if not ('faixa' in df.columns and 'partidas_faixa' in df.columns):
    mapa_faixa = {}
    for f in FAIXAS_PADRAO:
        for c in df.columns:
            if re.sub(r'[\s_]+', '', c) == f:
                mapa_faixa[c] = f
                break
    if not mapa_faixa:
        raise ValueError(
            "Tabela de entrada não está nem no formato longo "
            "(colunas 'faixa' e 'partidas_faixa') nem no formato wide "
            f"(colunas uma por faixa, ex.: '{next(iter(FAIXAS_PADRAO))}'). "
            f"Colunas lidas: {list(df.columns)}"
        )
    df = df.rename(columns=mapa_faixa)
    faixas_ausentes = [f for f in FAIXAS_PADRAO if f not in set(mapa_faixa.values())]
    if faixas_ausentes:
        print(f"   Aviso: formato wide sem as faixas {faixas_ausentes} (tratadas como ausentes).")
    id_vars = [c for c in df.columns if c not in set(mapa_faixa.values())]
    df = df.melt(
        id_vars=id_vars,
        value_vars=list(mapa_faixa.values()),
        var_name='faixa',
        value_name='partidas_faixa',
    )
    print(f"   Formato detectado: wide (uma coluna por faixa) — {len(mapa_faixa)} faixas derretidas para o formato longo")

obrigatorias = ['servico', 'vista', 'sentido', 'faixa', 'partidas_faixa']
faltantes = [c for c in obrigatorias if c not in df.columns]
if faltantes:
    raise ValueError(f"Colunas obrigatórias ausentes na tabela de entrada: {faltantes}. Colunas lidas: {list(df.columns)}")

print(f"1. Entrada carregada: {len(df):,} linhas, {df['servico'].nunique()} serviços")

# 2. Validação básica das colunas de classificação
dominio_ruim = sorted(set(df['sentido'].dropna().unique()) - SENTIDOS_VALIDOS)
if dominio_ruim:
    raise ValueError(f"Valores de 'sentido' fora do domínio {{Ida, Volta, Circular}}: {dominio_ruim}")

faixas_ruins = sorted(set(df['faixa'].dropna().unique()) - set(FAIXAS_PADRAO))
if faixas_ruins:
    raise ValueError(f"Valores de 'faixa' fora das 14 faixas padrão: {faixas_ruins}")

duplicatas = df.duplicated(['servico', 'sentido', 'faixa'])
if duplicatas.any():
    ex = df[duplicatas][['servico', 'sentido', 'faixa']].head(10).to_dict('records')
    raise ValueError(f"{duplicatas.sum()} linhas duplicadas em (servico, sentido, faixa). Exemplos: {ex}")

# 3. Conversão da contagem de partidas
valores = df['partidas_faixa'].fillna('').astype(str).str.strip()
mascara_traco = ~valores.str.fullmatch(r'\d+')
n_tracos = int(mascara_traco.sum())
if n_tracos:
    if traco_ausente == 'erro':
        raise ValueError(f"{n_tracos} valores não numéricos na coluna partidas_faixa (traco_ausente='erro').")
    print(f"   Aviso: {n_tracos} valores '-'/'vazios' em partidas_faixa tratados como 0.")
    valores = valores.where(~mascara_traco, '0')

df['n_partidas'] = pd.to_numeric(valores, errors='coerce')
if df['n_partidas'].isna().any():
    ex = df[df['n_partidas'].isna()][['servico', 'sentido', 'faixa', 'partidas_faixa']].head(10).to_dict('records')
    raise ValueError(f"Contagens não conversíveis em partidas_faixa. Exemplos: {ex}")
df['n_partidas'] = df['n_partidas'].astype(int)
if (df['n_partidas'] < 0).any():
    raise ValueError("Contagens negativas em partidas_faixa.")

# 4. Filtro de linhas
if linhas_rodar_set:
    df = df[df['servico'].str.upper().isin(linhas_rodar_set)].copy()
    print(f"   Filtro de linhas aplicado: {df['servico'].nunique()} serviços restantes")

# 5. Geração das bandas por (servico, sentido)
bandas = []
grupos_pulados = 0
grupos_processados = 0
nomes_arquivos = {}

grupos = df.groupby(['servico', 'sentido'], sort=True)
for (servico, sentido), grupo in grupos:
    n_total = int(grupo['n_partidas'].sum())
    if n_total == 0:
        grupos_pulados += 1
        continue

    vista = grupo['vista'].iloc[0]
    if grupo['vista'].nunique() > 1:
        raise ValueError(f"Grupo ({servico}, {sentido}) tem mais de uma vista: {sorted(grupo['vista'].unique())}")

    try:
        headsign = derivar_headsign(vista, sentido)
    except ValueError as e:
        raise ValueError(f"Grupo ({servico}, {sentido}) com partidas>0: {e}")

    nome_arquivo = f"horarios_{sanitize_filename(servico)}_{sanitize_filename(headsign)}_{service_id}.csv"
    if nome_arquivo in nomes_arquivos:
        outro = nomes_arquivos[nome_arquivo]
        raise ValueError(f"Colisão de nome de arquivo: {nome_arquivo} gerado por ({servico}, {sentido}) e por {outro}")
    nomes_arquivos[nome_arquivo] = (servico, sentido)

    for _, linha in grupo.iterrows():
        n = int(linha['n_partidas'])
        if n == 0:
            continue
        ini, fim = FAIXAS_PADRAO[linha['faixa']]
        dur = fim - ini
        h0 = dur // n
        h = (h0 // granularidade_headway) * granularidade_headway
        if h == 0:
            h = h0
        start = ini
        end = start + n * h
        bandas.append({
            'arquivo': nome_arquivo,
            'servico': servico,
            'sentido': sentido,
            'headsign': headsign,
            'faixa': linha['faixa'],
            'n_partidas': n,
            '_start': start,
            '_end': end,
            'start_time': seconds_to_horario(start),
            'end_time': seconds_to_horario(end),
            'headway_secs': h,
        })
    grupos_processados += 1

df_bandas = pd.DataFrame(bandas)
print(f"   Grupos com partidas: {grupos_processados} | pulados (todas as faixas zeradas): {grupos_pulados}")
print(f"   Bandas geradas: {len(df_bandas):,} | partidas: {int(df_bandas['n_partidas'].sum()):,}")

# 6. Gate de verificação (reexpansão x entrada) — aborta antes de gravar
print("\n2. Verificação (gate)...")

falhas = []

contagem_saida = df_bandas.groupby(['servico', 'sentido', 'faixa'])['n_partidas'].sum()
contagem_entrada = df[df['n_partidas'] > 0].groupby(['servico', 'sentido', 'faixa'])['n_partidas'].sum()
comparacao = contagem_entrada.to_frame('entrada').join(contagem_saida.to_frame('saida'), how='outer').fillna(0)
divergentes = comparacao[comparacao['entrada'] != comparacao['saida']]
if len(divergentes):
    falhas.append(f"{len(divergentes)} células (servico, sentido, faixa) com contagem divergente entrada x saída")

vazamentos = df_bandas[df_bandas['_end'] > df_bandas['faixa'].map(lambda f: FAIXAS_PADRAO[f][1])]
if len(vazamentos):
    falhas.append(f"{len(vazamentos)} bandas com end_time além do fim da faixa")

headways_ruins = df_bandas[(df_bandas['headway_secs'] < 1) | (df_bandas['_end'] <= df_bandas['_start'])]
if len(headways_ruins):
    falhas.append(f"{len(headways_ruins)} bandas com headway < 1 ou end_time <= start_time")

df_bandas['_h0'] = df_bandas.apply(lambda r: (FAIXAS_PADRAO[r['faixa']][1] - FAIXAS_PADRAO[r['faixa']][0]) // r['n_partidas'], axis=1)
multiplos = df_bandas[(df_bandas['_h0'] >= granularidade_headway) & (df_bandas['headway_secs'] % granularidade_headway != 0)]
if len(multiplos):
    falhas.append(f"{len(multiplos)} bandas com headway não múltiplo de {granularidade_headway}s")

total_entrada = int(df['n_partidas'].sum())
total_saida = int(df_bandas['n_partidas'].sum())
if total_entrada != total_saida:
    falhas.append(f"total de partidas entrada ({total_entrada:,}) != saída ({total_saida:,})")

print(f"   Células conferidas (servico, sentido, faixa): {len(comparacao):,}")
print(f"   Bandas fora da faixa: {len(vazamentos)} | headways inválidos: {len(headways_ruins)}")
print(f"   Partidas entrada/saída: {total_entrada:,} / {total_saida:,}")

if falhas:
    print("\n!!! GATE FALHOU:")
    for f in falhas:
        print(f"    - {f}")
    raise SystemExit(1)
print("   Gate: OK (contagens preservadas, nenhuma banda vaza da faixa)")

# 7. Gravação dos arquivos de quadro de horário
pasta_saida.mkdir(parents=True, exist_ok=True)
print(f"\n3. Gravando arquivos em: {pasta_saida}")

colunas_finais = ['trip_id', 'trip_headsign', 'trip_short_name', 'start_time', 'end_time', 'headway_secs']
total_arquivos = 0

for nome_arquivo, grupo in df_bandas.groupby('arquivo', sort=True):
    df_arq = grupo.sort_values(by='_start').copy()
    df_saida = pd.DataFrame({
        'trip_id': '',
        'trip_headsign': df_arq['headsign'],
        'trip_short_name': df_arq['servico'],
        'start_time': df_arq['start_time'],
        'end_time': df_arq['end_time'],
        'headway_secs': df_arq['headway_secs'].astype(str),
    })[colunas_finais]
    df_saida.to_csv(pasta_saida / nome_arquivo, index=False, encoding='utf-8')
    total_arquivos += 1

# 8. Exportação Excel (opcional): uma aba por linha, um bloco por quadro
if gerar_excel:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    def nome_aba_valido(servico: str, usados: set) -> str:
        """Nome de aba Excel válido (<= 31 chars, sem caracteres proibidos, único)."""
        nome = sanitize_filename(str(servico))[:31] or "SEM_NOME"
        base = nome
        i = 2
        while nome.lower() in usados:
            sufixo = f"_{i}"
            nome = base[:31 - len(sufixo)] + sufixo
            i += 1
        usados.add(nome.lower())
        return nome

    ORDEM_SENTIDO = {'Ida': 0, 'Volta': 1, 'Circular': 2}

    wb = Workbook()
    wb.remove(wb.active)
    abas_usadas = set()
    total_partidas_excel = 0

    for servico, quadros in df_bandas.groupby('servico', sort=True):
        ws = wb.create_sheet(title=nome_aba_valido(servico, abas_usadas))
        linha_atual = 1
        grupos = sorted(
            ((chave, grupo) for chave, grupo in quadros.groupby(['sentido', 'headsign'], sort=True)),
            key=lambda kv: (ORDEM_SENTIDO.get(kv[0][0], 9), kv[0][1]),
        )
        for (sentido, headsign), bandas in grupos:
            # Cabeçalho do quadro: linha, trip_headsign e tipo dia
            for col, (rotulo, valor) in enumerate([
                ("Linha", str(servico)),
                ("Trip headsign", headsign),
                ("Tipo dia", tipo_dia_rotulo),
            ], start=1):
                ws.cell(row=linha_atual, column=(col - 1) * 2 + 1, value=rotulo).font = Font(bold=True)
                ws.cell(row=linha_atual, column=(col - 1) * 2 + 2, value=valor)
            linha_atual += 2

            # Cabeçalho da tabela: Partida | Horário | Intervalo
            for col, rotulo in enumerate(['Partida', 'Horário', 'Intervalo'], start=1):
                ws.cell(row=linha_atual, column=col, value=rotulo).font = Font(bold=True)
            linha_atual += 1

            # Expansão das bandas em partidas individuais (uma por linha)
            partidas = []
            for _, banda in bandas.sort_values('_start').iterrows():
                for k in range(int(banda['n_partidas'])):
                    partidas.append(int(banda['_start']) + k * int(banda['headway_secs']))
            partidas.sort()
            anterior = None
            for idx, seg in enumerate(partidas, start=1):
                ws.cell(row=linha_atual, column=1, value=idx)
                ws.cell(row=linha_atual, column=2, value=seconds_to_horario(seg))
                intervalo = "--" if anterior is None else seconds_to_horario(seg - anterior)
                ws.cell(row=linha_atual, column=3, value=intervalo)
                linha_atual += 1
                anterior = seg
            total_partidas_excel += len(partidas)
            linha_atual += 1  # linha em branco entre blocos de quadro

        # Larguras de coluna: A-C tabela, D/E/F cabeçalho dos quadros
        for col in ('A', 'B', 'C', 'E', 'F'):
            ws.column_dimensions[col].width = 13
        ws.column_dimensions['D'].width = 28

    if total_partidas_excel != total_saida:
        raise SystemExit(
            f"!!! Excel: total expandido ({total_partidas_excel:,}) != partidas de saída ({total_saida:,})"
        )

    caminho_excel = pasta_saida / nome_excel
    wb.save(caminho_excel)
    excel_resumo = f"{len(wb.sheetnames)} abas, {total_partidas_excel:,} partidas"
    print(f"   Excel: {excel_resumo} -> {nome_excel}")
else:
    excel_resumo = "desativado"

# 9. Relatórios
if gerar_relatorio:
    relatorio = df_bandas.sort_values(['servico', 'sentido', '_start'])[
        ['arquivo', 'servico', 'sentido', 'headsign', 'faixa', 'n_partidas',
         'start_time', 'end_time', 'headway_secs']]
    relatorio.to_csv(pasta_saida / 'relatorio_bandas.csv', index=False, encoding='utf-8')

    resumo = [
        "RESUMO — 1.3 QH a partir de contagens de partidas",
        f"Tabela de entrada       : {caminho_partidas}",
        f"Linhas na tabela        : {len(df):,}",
        f"Grupos (servico,sentido): {grupos_processados} processados, {grupos_pulados} pulados (0 partidas)",
        f"Arquivos gerados        : {total_arquivos}",
        f"Bandas geradas          : {len(df_bandas):,}",
        f"Partidas entrada/saída  : {total_entrada:,} / {total_saida:,}",
        f"service_id              : {service_id}",
        f"Granularidade headway   : {granularidade_headway}s",
        f"Excel                   : {excel_resumo}",
        f"Gate                    : OK",
    ]
    (pasta_saida / '_resumo.txt').write_text('\n'.join(resumo) + '\n', encoding='utf-8')

print("\n==========================================================")
print("PROCESSAMENTO CONCLUÍDO COM SUCESSO!")
print(f"Arquivos de quadro gerados : {total_arquivos}")
print(f"Bandas                     : {len(df_bandas):,}")
print(f"Partidas (entrada = saída) : {total_saida:,}")
print(f"Excel                      : {excel_resumo}")
print(f"Pasta de destino           : {pasta_saida}")
print("==========================================================\n")
