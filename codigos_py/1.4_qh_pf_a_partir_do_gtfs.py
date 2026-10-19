import os
import re
import json
import math
import zipfile
import argparse
from pathlib import Path
import pandas as pd

# ==============================================================================
# LEITURA DE CONFIGURAÇÃO (Linha de Comando / Streamlit / JSON)
# ==============================================================================
_parser = argparse.ArgumentParser(
    description="Gera quadros de ponto facultativo (F_REG, formato do script 1) a partir das frequências de um GTFS, "
                "excluindo a MOBI-Rio e route_type != 700, com base DU (U_REG) nos consórcios e sábado (S_REG) em GTU/TUSE."
)
_parser.add_argument('--config', type=str, default='', help='Caminho para arquivo JSON de configuração')
_args, _ = _parser.parse_known_args()
_config = {}
if _args.config:
    if os.path.isfile(_args.config):
        with open(_args.config, 'r', encoding='utf-8-sig') as _f:
            _config = json.load(_f)
    else:
        try:
            _config = json.loads(_args.config)
        except Exception:
            with open(_args.config, 'r', encoding='utf-8-sig') as _f:
                _config = json.load(_f)

# Parâmetros principais
BASE_RESULTADOS = Path(_config.get("BASE_RESULTADOS", "C:/R_SMTR/resultados"))

ano_gtfs = str(_config.get("ano_gtfs", "2026"))
mes_gtfs = str(_config.get("mes_gtfs", "10")).zfill(2)
estudo_gtfs = str(_config.get("estudo_gtfs", "02"))

# Caminho do GTFS (ZIP na raiz do repositório por default)
_REPO_RAIZ = Path(__file__).resolve().parent.parent
end_gtfs = Path(_config.get("end_gtfs", _REPO_RAIZ / "0171_gtfs_out_26_1E_ret.zip"))

# service_id dos arquivos gerados (ponto facultativo)
service_id = str(_config.get("service_id", "F_REG"))

# Agências excluídas integralmente (MOBI-Rio em qualquer route_type)
agencias_excluidas = _config.get("agencias_excluidas", ["20001"])
if isinstance(agencias_excluidas, str):
    agencias_excluidas = [x.strip() for x in agencias_excluidas.split(',') if x.strip()]
agencias_excluidas = {str(x) for x in agencias_excluidas}

# route_types permitidos (apenas ônibus convencional)
route_types_permitidos = _config.get("route_types_permitidos", ["700"])
if isinstance(route_types_permitidos, str):
    route_types_permitidos = [x.strip() for x in route_types_permitidos.split(',') if x.strip()]
route_types_permitidos = {str(x) for x in route_types_permitidos}

# Base de partidas por agência: service_id de origem no GTFS
#   DU (U_REG): Internorte, Intersul, Transcarioca, Santa Cruz
#   sábado (S_REG): GTU, TUSE
mapa_base_default = {
    "22003": "U_REG",  # Internorte
    "22002": "U_REG",  # Intersul
    "22004": "U_REG",  # Transcarioca
    "22005": "U_REG",  # Santa Cruz
    "A2": "S_REG",     # GTU
    "B2": "S_REG",     # TUSE
}
mapa_base = {str(k): str(v) for k, v in _config.get("mapa_base", mapa_base_default).items()}

# Linhas a rodar (vazio ou ['todas'] = todas)
linhas_rodar = _config.get("linhas_rodar", [])
if isinstance(linhas_rodar, str):
    linhas_rodar = [x.strip() for x in linhas_rodar.split(',') if x.strip()]
linhas_rodar_set = {str(x).upper() for x in linhas_rodar} if linhas_rodar and ('todas' not in [str(x).lower() for x in linhas_rodar]) else set()

# Pasta de saída (pasta própria; não mistura com o 1.3)
pasta_saida_config = _config.get("pasta_saida", "")
if pasta_saida_config:
    pasta_saida = Path(pasta_saida_config)
else:
    pasta_saida = BASE_RESULTADOS / f"quadro_horario_extraido/{ano_gtfs}/{mes_gtfs}/qh_por_linha_pf_gtfs/{estudo_gtfs}Q"

gerar_relatorio = bool(_config.get("gerar_relatorio", True))

# Exportação Excel: um workbook com uma aba por linha;
# em cada aba, um bloco por quadro (headsign) com Partida | Horário | Intervalo.
gerar_excel = bool(_config.get("gerar_excel", True))
nome_excel = str(_config.get("nome_excel", f"quadros_partidas_{service_id}.xlsx"))
tipo_dia_rotulo = str(_config.get("tipo_dia", service_id))

# ==============================================================================
# CONSTANTES E FUNÇÕES UTILITÁRIAS
# ==============================================================================
def sanitize_filename(name: str) -> str:
    """Substitui caracteres inválidos para nomes de arquivos no sistema operacional."""
    if not name:
        return "SEM_NOME"
    sanitized = re.sub(r'[\\/*?:"<>|]', '_', str(name))
    sanitized = re.sub(r'\s+', ' ', sanitized).strip('. ')
    return sanitized if sanitized else "SEM_NOME"


def seconds_to_horario(total_seconds: int) -> str:
    """Converte total de segundos para string HH:MM:SS mantendo H >= 24."""
    if total_seconds < 0:
        return ""
    h = total_seconds // 3600
    m = (total_seconds % 3600) // 60
    s = total_seconds % 60
    return f"{h:02d}:{m:02d}:{s:02d}"


def horario_to_seconds(hk_str) -> int:
    """Converte horário HH:MM:SS ou HH:MM para total de segundos (suporta H >= 24)."""
    if pd.isna(hk_str) or not str(hk_str).strip():
        return -1
    parts = str(hk_str).strip().split(':')
    if len(parts) < 2:
        return -1
    try:
        h = int(parts[0])
        m = int(parts[1])
        s = int(parts[2]) if len(parts) >= 3 else 0
        return h * 3600 + m * 60 + s
    except (ValueError, TypeError):
        return -1


def n_partidas_banda(start: int, end: int, headway: int) -> int:
    """Partidas de uma banda de frequência (exact_times=0): start, start+h, ... enquanto < end."""
    if headway < 1 or end <= start:
        return 0
    return math.ceil((end - start) / headway)


# ==============================================================================
# EXECUÇÃO PRINCIPAL
# ==============================================================================
print("==========================================================")
print("GERADOR DE QUADROS DE PONTO FACULTATIVO A PARTIR DO GTFS")
print(f"GTFS de Entrada   : {end_gtfs}")
print(f"Pasta de Destino  : {pasta_saida}")
print(f"service_id saída  : {service_id}")
print(f"route_types       : {sorted(route_types_permitidos)}")
print(f"Agências excluídas: {sorted(agencias_excluidas)}")
print("==========================================================\n")

if not end_gtfs.exists():
    raise FileNotFoundError(f"GTFS não encontrado: {end_gtfs}")

# 1. Leitura do GTFS (ZIP)
with zipfile.ZipFile(end_gtfs, 'r') as z:
    frequencies = pd.read_csv(z.open('frequencies.txt'), dtype=str)
    trips = pd.read_csv(z.open('trips.txt'), dtype=str)
    routes = pd.read_csv(z.open('routes.txt'), dtype=str)
    agency = pd.read_csv(z.open('agency.txt'), dtype=str)

print(f"1. Frequências lidas: {len(frequencies):,} | trips: {len(trips):,} | rotas: {len(routes):,}")

# 2. Merge e filtros
df = (
    frequencies
    .merge(trips[['trip_id', 'route_id', 'service_id', 'trip_headsign', 'trip_short_name']], on='trip_id', how='left')
    .merge(routes[['route_id', 'agency_id', 'route_type', 'route_short_name']], on='route_id', how='left')
    .merge(agency[['agency_id', 'agency_name']], on='agency_id', how='left')
)

n_antes = len(df)
df = df[df['agency_id'].isin(mapa_base)].copy()
n_sem_agencia = n_antes - len(df)

df = df[~df['agency_id'].isin(agencias_excluidas)].copy()
df = df[df['route_type'].isin(route_types_permitidos)].copy()
df = df[df['service_id'] == df['agency_id'].map(mapa_base)].copy()
print(f"   Bandas após filtros (agências base, sem MOBI, route_type {sorted(route_types_permitidos)}, service DU/SAB): {len(df):,}")

if df.empty:
    raise SystemExit("Nenhuma banda de frequência sobrou após os filtros.")

# 3. Filtro de linhas (opcional)
if linhas_rodar_set:
    df = df[df['trip_short_name'].astype(str).str.upper().isin(linhas_rodar_set)].copy()
    print(f"   Filtro de linhas aplicado: {len(df):,} bandas restantes")

# 4. Avisos de rotas sem bandas na base esperada
rotas_esperadas = (
    routes[
        routes['agency_id'].isin(mapa_base) &
        (~routes['agency_id'].isin(agencias_excluidas)) &
        (routes['route_type'].isin(route_types_permitidos))
    ]
    .merge(agency[['agency_id', 'agency_name']], on='agency_id', how='left')
)
rotas_com_banda = set(df['route_id'])
sem_banda = sorted(set(rotas_esperadas['route_id']) - rotas_com_banda)
if sem_banda:
    nomes = rotas_esperadas[rotas_esperadas['route_id'].isin(sem_banda)][['route_short_name', 'agency_name']]
    lista = ', '.join(f"{r} ({a})" for r, a in nomes.itertuples(index=False))
    print(f"   Aviso: {len(sem_banda)} rota(s) sem bandas na base de partidas: {lista}")

# 5. Colunas de tempo em segundos
df['_start'] = df['start_time'].map(horario_to_seconds)
df['_end'] = df['end_time'].map(horario_to_seconds)
df['_headway'] = pd.to_numeric(df['headway_secs'], errors='coerce')
df['service_id_origem'] = df['service_id']

tempo_ruino = df[(df['_start'] < 0) | (df['_end'] < 0) | df['_headway'].isna()]
if len(tempo_ruino):
    ex = tempo_ruino[['trip_id', 'start_time', 'end_time', 'headway_secs']].head(10).to_dict('records')
    raise ValueError(f"{len(tempo_ruino)} bandas com horário/headway inválidos. Exemplos: {ex}")
df['_headway'] = df['_headway'].astype(int)

# 6. Geração dos quadros por (trip_short_name, trip_headsign)
df['headsign'] = df['trip_headsign'].fillna('SEM_HEADSIGN').astype(str).str.strip()
df['headsign'] = df['headsign'].replace('', 'SEM_HEADSIGN')

bandas = []
nomes_arquivos = {}
grupos = df.groupby(['trip_short_name', 'trip_headsign'], sort=True)
for (servico, vista), grupo in grupos:
    headsign = grupo['headsign'].iloc[0]
    nome_arquivo = f"horarios_{sanitize_filename(servico)}_{sanitize_filename(headsign)}_{service_id}.csv"
    if nome_arquivo in nomes_arquivos:
        outro = nomes_arquivos[nome_arquivo]
        raise ValueError(f"Colisão de nome de arquivo: {nome_arquivo} gerado por ({servico}, {vista}) e por {outro}")
    nomes_arquivos[nome_arquivo] = (servico, vista)

    for _, linha in grupo.iterrows():
        bandas.append({
            'arquivo': nome_arquivo,
            'servico': servico,
            'headsign': headsign,
            'service_id_origem': linha['service_id_origem'],
            '_start': int(linha['_start']),
            '_end': int(linha['_end']),
            '_headway': int(linha['_headway']),
            'n_partidas': n_partidas_banda(int(linha['_start']), int(linha['_end']), int(linha['_headway'])),
        })

df_bandas = pd.DataFrame(bandas)
print(f"2. Quadros: {df_bandas['arquivo'].nunique()} | bandas geradas: {len(df_bandas):,} | partidas expandidas: {int(df_bandas['n_partidas'].sum()):,}")

# 7. Gate de verificação (antes de gravar)
print("\n3. Verificação (gate)...")

falhas = []

bandas_ruins = df_bandas[(df_bandas['_headway'] < 1) | (df_bandas['_end'] <= df_bandas['_start'])]
if len(bandas_ruins):
    falhas.append(f"{len(bandas_ruins)} bandas com headway < 1 ou end <= start")

partidas_zero = df_bandas[df_bandas['n_partidas'] < 1]
if len(partidas_zero):
    falhas.append(f"{len(partidas_zero)} bandas com 0 partidas na reexpansão")

# preservação entrada x saída por (linha, headsign, start_time, end_time, headway)
chaves = ['servico', 'headsign', '_start', '_end', '_headway']
entrada = df.groupby(['trip_short_name', 'headsign', '_start', '_end', '_headway']).size().to_frame('entrada')
entrada.index.names = chaves
saida = df_bandas.groupby(chaves).size().to_frame('saida')
comparacao = entrada.join(saida, how='outer').fillna(0)
divergentes = comparacao[comparacao['entrada'] != comparacao['saida']]
if len(divergentes):
    falhas.append(f"{len(divergentes)} bandas (linha, headsign, start, end, headway) divergentes entrada x saída")

partidas_entradas_calc = int(sum(
    n_partidas_banda(int(s), int(e), int(h))
    for s, e, h in zip(df['_start'], df['_end'], df['_headway'])
))
partidas_saida = int(df_bandas['n_partidas'].sum())
if partidas_entradas_calc != partidas_saida:
    falhas.append(f"total de partidas entrada ({partidas_entradas_calc:,}) != saída ({partidas_saida:,})")

print(f"   Bandas conferidas: {len(comparacao):,} | divergências: {len(divergentes)}")
print(f"   Partidas entrada/saída: {partidas_entradas_calc:,} / {partidas_saida:,}")

if falhas:
    print("\n!!! GATE FALHOU:")
    for f in falhas:
        print(f"    - {f}")
    raise SystemExit(1)
print("   Gate: OK (bandas e partidas preservadas)")

# 8. Gravação dos CSVs de quadro de horário
pasta_saida.mkdir(parents=True, exist_ok=True)
print(f"\n4. Gravando arquivos em: {pasta_saida}")

colunas_finais = ['trip_id', 'trip_headsign', 'trip_short_name', 'start_time', 'end_time', 'headway_secs']
total_arquivos = 0

for nome_arquivo, grupo in df_bandas.groupby('arquivo', sort=True):
    df_saida = pd.DataFrame({
        'trip_id': '',
        'trip_headsign': grupo['headsign'],
        'trip_short_name': grupo['servico'],
        'start_time': grupo['_start'].map(seconds_to_horario),
        'end_time': grupo['_end'].map(seconds_to_horario),
        'headway_secs': grupo['_headway'].astype(str),
    })[colunas_finais].sort_values('start_time')
    df_saida.to_csv(pasta_saida / nome_arquivo, index=False, encoding='utf-8')
    total_arquivos += 1

# 9. Exportação Excel (opcional): uma aba por linha, um bloco por quadro
if gerar_excel:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    def nome_aba_valido(servico: str, usados: set) -> str:
        """Nome de aba Excel válido (<= 31 chars, sem caracteres proibidos, único)."""
        nome = sanitize_filename(str(servico))[:31] or "SEM_NOME"
        base = nome
        i = 2
        while nome.lower() in usados:
            sufixo = f"_{i}"
            nome = base[:31 - len(sufixo)] + sufixo
            i += 1
        usados.add(nome.lower())
        return nome

    wb = Workbook()
    wb.remove(wb.active)
    abas_usadas = set()
    total_partidas_excel = 0

    for servico, quadros in df_bandas.groupby('servico', sort=True):
        ws = wb.create_sheet(title=nome_aba_valido(servico, abas_usadas))
        linha_atual = 1
        for headsign, bandas_linha in sorted(quadros.groupby('headsign', sort=True)):
            for col, (rotulo, valor) in enumerate([
                ("Linha", str(servico)),
                ("Trip headsign", headsign),
                ("Tipo dia", tipo_dia_rotulo),
            ], start=1):
                ws.cell(row=linha_atual, column=(col - 1) * 2 + 1, value=rotulo).font = Font(bold=True)
                ws.cell(row=linha_atual, column=(col - 1) * 2 + 2, value=valor)
            linha_atual += 2

            for col, rotulo in enumerate(['Partida', 'Horário', 'Intervalo'], start=1):
                ws.cell(row=linha_atual, column=col, value=rotulo).font = Font(bold=True)
            linha_atual += 1

            partidas = []
            for _, banda in bandas_linha.sort_values('_start').iterrows():
                for k in range(int(banda['n_partidas'])):
                    partidas.append(int(banda['_start']) + k * int(banda['_headway']))
            partidas.sort()
            anterior = None
            for idx, seg in enumerate(partidas, start=1):
                ws.cell(row=linha_atual, column=1, value=idx)
                ws.cell(row=linha_atual, column=2, value=seconds_to_horario(seg))
                intervalo = "--" if anterior is None else seconds_to_horario(seg - anterior)
                ws.cell(row=linha_atual, column=3, value=intervalo)
                linha_atual += 1
                anterior = seg
            total_partidas_excel += len(partidas)
            linha_atual += 1

        for col in ('A', 'B', 'C', 'E', 'F'):
            ws.column_dimensions[col].width = 13
        ws.column_dimensions['D'].width = 28

    if total_partidas_excel != partidas_saida:
        raise SystemExit(
            f"!!! Excel: total expandido ({total_partidas_excel:,}) != partidas de saída ({partidas_saida:,})"
        )

    caminho_excel = pasta_saida / nome_excel
    wb.save(caminho_excel)
    excel_resumo = f"{len(wb.sheetnames)} abas, {total_partidas_excel:,} partidas"
    print(f"   Excel: {excel_resumo} -> {nome_excel}")
else:
    excel_resumo = "desativado"

# 10. Relatórios
if gerar_relatorio:
    relatorio = df_bandas.sort_values(['servico', 'headsign', '_start']).copy()
    relatorio['start_time'] = relatorio['_start'].map(seconds_to_horario)
    relatorio['end_time'] = relatorio['_end'].map(seconds_to_horario)
    relatorio['headway_secs'] = relatorio['_headway']
    relatorio[
        ['arquivo', 'servico', 'headsign', 'service_id_origem', 'start_time', 'end_time', 'headway_secs', 'n_partidas']
    ].to_csv(pasta_saida / 'relatorio_bandas.csv', index=False, encoding='utf-8')

    por_origem = df_bandas.groupby('service_id_origem').agg(
        quadros=('arquivo', 'nunique'), bandas=('servico', 'size'), partidas=('n_partidas', 'sum'))
    resumo = [
        "RESUMO — 1.4 QH ponto facultativo a partir do GTFS",
        f"GTFS de entrada        : {end_gtfs}",
        f"Pasta de destino       : {pasta_saida}",
        f"Rotas sem bandas       : {len(sem_banda)}" + (f" ({lista})" if sem_banda else ""),
        f"Quadros gerados        : {total_arquivos}",
        f"Bandas geradas         : {len(df_bandas):,}",
        f"Partidas (entrada=saída): {partidas_saida:,}",
        f"service_id             : {service_id}",
        f"Excel                  : {excel_resumo}",
        f"Gate                   : OK",
        "",
        "Por base de partidas (service_id_origem):",
        por_origem.to_string(),
    ]
    (pasta_saida / '_resumo.txt').write_text('\n'.join(resumo) + '\n', encoding='utf-8')

print("\n==========================================================")
print("PROCESSAMENTO CONCLUÍDO COM SUCESSO!")
print(f"Arquivos de quadro gerados : {total_arquivos}")
print(f"Bandas                     : {len(df_bandas):,}")
print(f"Partidas (entrada = saída) : {partidas_saida:,}")
print(f"Excel                      : {excel_resumo}")
print(f"Pasta de destino           : {pasta_saida}")
print("==========================================================\n")
